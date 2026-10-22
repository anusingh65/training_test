#!/usr/bin/env python3
"""Generate a PowerPoint presentation matching presentation.html exactly."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ── Colors matching the HTML CSS variables ──
PRIMARY = RGBColor(0x0A, 0x0F, 0x1A)
SECONDARY = RGBColor(0x11, 0x18, 0x27)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
ACCENT_LIGHT = RGBColor(0x3B, 0x82, 0xF6)
PURPLE = RGBColor(0x7C, 0x3A, 0xED)
PURPLE_LIGHT = RGBColor(0x8B, 0x5C, 0xF6)
TEXT_COLOR = RGBColor(0xF8, 0xFA, 0xFC)
TEXT_MUTED = RGBColor(0x94, 0xA3, 0xB8)
SURFACE = RGBColor(0x1E, 0x29, 0x3B)
GREEN = RGBColor(0x10, 0xB9, 0x81)
RED = RGBColor(0xEF, 0x44, 0x44)
YELLOW = RGBColor(0xF5, 0x9E, 0x0B)
ORANGE = RGBColor(0xFB, 0x92, 0x3C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_slide_bg(slide, color):
    """Set slide background to a solid color."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=TEXT_COLOR, bold=False, italic=False, alignment=PP_ALIGN.LEFT,
                font_name='Calibri'):
    """Add a simple textbox with one run of text."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font_name
    return txBox, tf, p


def add_multirun_paragraph(tf, runs, alignment=PP_ALIGN.LEFT, space_after=Pt(0)):
    """Add a paragraph with multiple runs of different formatting.
    runs = [(text, font_size, color, bold, italic, font_name), ...]
    """
    p = tf.add_paragraph() if len(tf.paragraphs) > 0 and tf.paragraphs[0].text != '' else tf.paragraphs[0]
    if len(tf.paragraphs) > 1 or tf.paragraphs[0].text != '':
        p = tf.add_paragraph()
    p.alignment = alignment
    p.space_after = space_after
    for text, font_size, color, bold, italic, font_name in runs:
        run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font_name or 'Calibri'
    return p


def add_rounded_rect(slide, left, top, width, height, fill_color=None,
                     border_color=None, border_width=Pt(1)):
    """Add a rounded rectangle shape."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width
        shape.line.fill.solid()
        shape.line.fill.fore_color.rgb = border_color
    else:
        shape.line.fill.background()
    # Reduce corner rounding
    shape.adjustments[0] = 0.05
    return shape


def add_circle(slide, left, top, size, fill_color):
    """Add a circle (oval) shape."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, left, top, size, size
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape


def add_rect(slide, left, top, width, height, fill_color):
    """Add a plain rectangle shape."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape


def add_slide_number(slide, number):
    """Add slide number in bottom-right corner."""
    add_textbox(slide, Inches(11.5), Inches(7.0), Inches(1), Inches(0.4),
                f"{number:02d}", font_size=12, color=TEXT_MUTED,
                alignment=PP_ALIGN.RIGHT, font_name='Calibri')


def add_section_label(slide, text, left=Inches(0.8), top=Inches(0.6)):
    """Add the small uppercase section label (h3 equivalent)."""
    add_textbox(slide, left, top, Inches(4), Inches(0.4),
                text.upper(), font_size=11, color=ACCENT, bold=True,
                font_name='Calibri')


def add_section_title(slide, text, left=Inches(0.8), top=Inches(1.0), width=Inches(10)):
    """Add the large section title (h2 equivalent)."""
    add_textbox(slide, left, top, width, Inches(0.8),
                text, font_size=36, color=TEXT_COLOR, bold=False,
                font_name='Georgia')


def add_body_text(slide, text, left=Inches(0.8), top=Inches(1.9), width=Inches(10)):
    """Add body paragraph text."""
    txBox = slide.shapes.add_textbox(left, top, width, Inches(1.0))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(16)
    run.font.color.rgb = TEXT_MUTED
    run.font.name = 'Calibri'
    return txBox


def add_stat_block(slide, left, top, number, label, number_color=ACCENT):
    """Add a stat with large number and small label below it."""
    add_textbox(slide, left, top, Inches(2.2), Inches(0.7),
                number, font_size=44, color=number_color,
                alignment=PP_ALIGN.CENTER, font_name='Georgia')
    add_textbox(slide, left, top + Inches(0.65), Inches(2.2), Inches(0.4),
                label.upper(), font_size=9, color=TEXT_MUTED, bold=True,
                alignment=PP_ALIGN.CENTER)


def add_card_with_border(slide, left, top, width, height, title, body,
                         border_side='left', border_color=ACCENT,
                         title_color=TEXT_COLOR, body_color=TEXT_MUTED):
    """Add a card with colored left/top border, title and body text."""
    # Background card
    bg = add_rounded_rect(slide, left, top, width, height,
                          fill_color=SURFACE, border_color=RGBColor(0x2A, 0x35, 0x4B))
    # Border accent line
    if border_side == 'left':
        add_rect(slide, left, top + Emu(Inches(0.08).emu), Pt(3), height - Emu(Inches(0.16).emu), border_color)
    elif border_side == 'top':
        add_rect(slide, left + Emu(Inches(0.08).emu), top, width - Emu(Inches(0.16).emu), Pt(3), border_color)
    # Title
    add_textbox(slide, left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.35),
                title, font_size=13, color=title_color, bold=True)
    # Body
    txBox = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.48),
                                     width - Inches(0.4), height - Inches(0.6))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = body
    run.font.size = Pt(10)
    run.font.color.rgb = body_color
    run.font.name = 'Calibri'


def add_rec_item(slide, left, top, width, title, body, border_color=ACCENT, title_color=ACCENT_LIGHT):
    """Add a recommendation-style item with left border."""
    height = Inches(0.85)
    # Background
    bg = add_rounded_rect(slide, left, top, width, height,
                          fill_color=RGBColor(0x14, 0x1C, 0x2E),
                          border_color=RGBColor(0x1E, 0x29, 0x3B))
    # Left border
    add_rect(slide, left, top + Emu(Inches(0.05).emu), Pt(3),
             height - Emu(Inches(0.1).emu), border_color)
    # Title
    add_textbox(slide, left + Inches(0.18), top + Inches(0.06), width - Inches(0.3), Inches(0.28),
                title, font_size=11, color=title_color, bold=True)
    # Body
    txBox = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.32),
                                     width - Inches(0.3), height - Inches(0.38))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = body
    run.font.size = Pt(9)
    run.font.color.rgb = TEXT_MUTED
    run.font.name = 'Calibri'
    return height


# ══════════════════════════════════════════════════════════════
# Create Presentation  (16:9 widescreen)
# ══════════════════════════════════════════════════════════════
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height


# ────────────────────────── SLIDE 1: TITLE ──────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
set_slide_bg(slide, PRIMARY)

# Team name
add_textbox(slide, Inches(0), Inches(0.8), SLIDE_W, Inches(0.4),
            'TEAM CODE DIGGERS', font_size=11, color=ACCENT, bold=True,
            alignment=PP_ALIGN.CENTER)

# Title
add_textbox(slide, Inches(0), Inches(1.3), SLIDE_W, Inches(1.2),
            'MediChain', font_size=72, color=TEXT_COLOR, bold=False,
            alignment=PP_ALIGN.CENTER, font_name='Georgia')

# Tagline
add_textbox(slide, Inches(1.5), Inches(2.5), Inches(10.333), Inches(0.6),
            'Blockchain + AI System for Counterfeit Drug Prevention',
            font_size=22, color=TEXT_MUTED, italic=True,
            alignment=PP_ALIGN.CENTER)

# Problem statement quote box
quote_left = Inches(2.2)
quote_top = Inches(3.2)
quote_w = Inches(8.9)
quote_h = Inches(1.0)
bg = add_rounded_rect(slide, quote_left, quote_top, quote_w, quote_h,
                       fill_color=RGBColor(0x0E, 0x15, 0x26),
                       border_color=RGBColor(0x1A, 0x25, 0x3A))
add_rect(slide, quote_left, quote_top + Emu(Inches(0.05).emu), Pt(3),
         quote_h - Emu(Inches(0.1).emu), ACCENT)

txBox = slide.shapes.add_textbox(quote_left + Inches(0.3), quote_top + Inches(0.12),
                                 quote_w - Inches(0.6), quote_h - Inches(0.2))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
run = p.add_run()
run.text = '\u201cDesign a blockchain-based drug traceability platform where each medicine strip is verified through a unique QR code, with AI models analyzing scan patterns to predict high-risk counterfeiting zones.\u201d'
run.font.size = Pt(12)
run.font.color.rgb = TEXT_MUTED
run.font.italic = True
run.font.name = 'Calibri'

# Stats
stat_y = Inches(4.7)
add_stat_block(slide, Inches(2.3), stat_y, '1M+', 'Deaths per Year', RED)
add_stat_block(slide, Inches(5.5), stat_y, '$200B', 'Lost Annually', YELLOW)
add_stat_block(slide, Inches(8.7), stat_y, '10\u201330%', 'Drugs Are Fake (WHO)', ORANGE)

add_slide_number(slide, 1)


# ────────────────────────── SLIDE 2: THE CRISIS ──────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, PRIMARY)
add_section_label(slide, 'The Crisis')
add_section_title(slide, 'The System is Broken')
add_body_text(slide, 'India is the world\u2019s 3rd largest pharmaceutical producer \u2014 yet $4.1B is lost annually to counterfeits. There are zero real-time tracking systems. Patients die. Pharmacies lose trust. Criminals go undetected.')

# Stats row
stat_y = Inches(2.8)
add_stat_block(slide, Inches(1.0), stat_y, '3rd', 'Largest Pharma Producer', ACCENT)
add_stat_block(slide, Inches(4.5), stat_y, '$4.1B', 'Lost to Counterfeits in India', RED)
add_stat_block(slide, Inches(8.0), stat_y, '0', 'Real-Time Tracking Systems', YELLOW)

# Problem cards (2x2)
problems = [
    ('Opaque', 'No visibility into where a drug has been or who handled it.'),
    ('Reactive', 'Counterfeits discovered only AFTER patients are harmed.'),
    ('No Prediction', 'Cannot identify risk zones before outbreaks happen.'),
    ('Paper-Based', 'Documents easily forged. Impossible to audit at scale.'),
]
card_w = Inches(5.5)
card_h = Inches(0.7)
positions = [
    (Inches(0.8), Inches(4.3)), (Inches(6.8), Inches(4.3)),
    (Inches(0.8), Inches(5.2)), (Inches(6.8), Inches(5.2)),
]
for (title, body), (lx, ly) in zip(problems, positions):
    bg = add_rounded_rect(slide, lx, ly, card_w, card_h,
                          fill_color=RGBColor(0x14, 0x1C, 0x2E),
                          border_color=RGBColor(0x1E, 0x29, 0x3B))
    # Red dot indicator
    add_circle(slide, lx + Inches(0.2), ly + Inches(0.18), Inches(0.3), RGBColor(0x2D, 0x13, 0x15))
    add_textbox(slide, lx + Inches(0.23), ly + Inches(0.18), Inches(0.3), Inches(0.3),
                '\u2716', font_size=11, color=RED, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, lx + Inches(0.6), ly + Inches(0.08), card_w - Inches(0.8), Inches(0.3),
                title, font_size=13, color=TEXT_COLOR, bold=True)
    add_textbox(slide, lx + Inches(0.6), ly + Inches(0.35), card_w - Inches(0.8), Inches(0.3),
                body, font_size=10, color=TEXT_MUTED)

add_slide_number(slide, 2)


# ────────────────────────── SLIDE 3: SOLUTION OVERVIEW ──────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, PRIMARY)
add_section_label(slide, 'Our Approach')
add_section_title(slide, 'Not Just QR Codes')
add_body_text(slide, 'Every team will put a QR code on a pill. That\u2019s the baseline. The real question: what happens when the counterfeiter is smarter than your QR code?')

# Three step cards
steps = [
    ('1', 'Blockchain Layer', 'Every medicine registered on Ethereum. Every transfer recorded permanently. 3 smart contracts enforce supply chain rules.', ACCENT),
    ('2', 'AI Layer', 'Unsupervised anomaly detection catches counterfeits without needing labeled examples. Predictive heatmap forecasts risk zones.', PURPLE),
    ('3', 'Government Loop', 'Heatmap feeds CDSCO in real-time. Raid data feeds back as ground truth. The system gets smarter with every enforcement action.', GREEN),
]
card_w = Inches(3.6)
card_h = Inches(2.2)
start_x = Inches(1.2)
gap = Inches(0.4)
card_y = Inches(3.0)

for i, (num, title, body, color) in enumerate(steps):
    lx = start_x + i * (card_w + gap)
    bg = add_rounded_rect(slide, lx, card_y, card_w, card_h,
                          fill_color=RGBColor(0x12, 0x1A, 0x2B),
                          border_color=RGBColor(0x1E, 0x29, 0x3B))
    # Top border
    add_rect(slide, lx + Emu(Inches(0.05).emu), card_y, card_w - Emu(Inches(0.1).emu), Pt(3), color)
    # Number circle
    circle = add_circle(slide, lx + Inches(1.45), card_y + Inches(0.25), Inches(0.45), ACCENT)
    add_textbox(slide, lx + Inches(1.45), card_y + Inches(0.27), Inches(0.45), Inches(0.45),
                num, font_size=16, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    # Title
    add_textbox(slide, lx + Inches(0.2), card_y + Inches(0.8), card_w - Inches(0.4), Inches(0.35),
                title, font_size=15, color=TEXT_COLOR, bold=True, alignment=PP_ALIGN.CENTER)
    # Body
    txBox = slide.shapes.add_textbox(lx + Inches(0.2), card_y + Inches(1.2),
                                     card_w - Inches(0.4), card_h - Inches(1.4))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = body
    run.font.size = Pt(10)
    run.font.color.rgb = TEXT_MUTED
    run.font.name = 'Calibri'

# Bottom callout box
callout_y = Inches(5.5)
callout_w = Inches(10)
callout_h = Inches(0.7)
callout_x = Inches(1.7)
bg = add_rounded_rect(slide, callout_x, callout_y, callout_w, callout_h,
                       fill_color=RGBColor(0x12, 0x14, 0x28),
                       border_color=RGBColor(0x2A, 0x20, 0x4B))
add_rect(slide, callout_x, callout_y + Emu(Inches(0.05).emu), Pt(3),
         callout_h - Emu(Inches(0.1).emu), PURPLE)
add_textbox(slide, callout_x + Inches(0.3), callout_y + Inches(0.12), callout_w - Inches(0.6), callout_h - Inches(0.2),
            'We don\u2019t just verify \u2014 we predict. Our AI tells you where counterfeits will appear next week, not where they were found yesterday.',
            font_size=12, color=TEXT_COLOR, bold=False, alignment=PP_ALIGN.LEFT)

add_slide_number(slide, 3)


# ────────────────────────── SLIDE 4: TECHNICAL ARCHITECTURE ──────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, PRIMARY)
add_section_label(slide, 'Under the Hood')
add_section_title(slide, 'Architecture & AI Models')

# ── LEFT COLUMN: Anomaly Detection ──
col_left = Inches(0.6)
col_top = Inches(1.8)
col_w = Inches(5.8)

add_textbox(slide, col_left, col_top, col_w, Inches(0.3),
            'AI MODEL 1 \u2014 ANOMALY DETECTION', font_size=9, color=ACCENT, bold=True)

items_left = [
    ('Isolation Forest', 'Unsupervised ML. Trains on normal supply chain data only. No labeled counterfeit examples needed \u2014 isolates anomalies by how few random splits it takes to separate them.', ACCENT, ACCENT_LIGHT),
    ('Autoencoder', 'Neural network trained to reconstruct normal transactions. When reconstruction error is high, the transaction doesn\u2019t match learned patterns \u2192 anomaly.', PURPLE, PURPLE_LIGHT),
    ('6 Domain Rules', 'Velocity > 500 km/h, path deviation, >5 transfers, 2\u20135am weekend timing, batch contamination >10%, low actor reputation.', GREEN, GREEN),
]

iy = col_top + Inches(0.35)
for title, body, border_col, title_col in items_left:
    add_rec_item(slide, col_left, iy, col_w, title, body,
                 border_color=border_col, title_color=title_col)
    iy += Inches(0.92)

# Feature pills label
add_textbox(slide, col_left, iy + Inches(0.05), col_w, Inches(0.25),
            '12 ENGINEERED FEATURES', font_size=9, color=PURPLE, bold=True)

# Feature pills
features = ['Velocity', 'Transit Time', 'Transfer Count', 'Path Deviation',
            'Scan Frequency', 'Unique Devices', 'Hour of Day', 'Is Weekend',
            'Days Since Mfg', 'Batch Flag Rate', 'Actor Reputation', 'Geo Consistency']
pill_y = iy + Inches(0.35)
pill_x = col_left
pill_h = Inches(0.28)
for feat in features:
    pw = Inches(len(feat) * 0.09 + 0.3)
    if pill_x + pw > col_left + col_w:
        pill_x = col_left
        pill_y += Inches(0.32)
    bg = add_rounded_rect(slide, pill_x, pill_y, pw, pill_h,
                          fill_color=RGBColor(0x0E, 0x17, 0x2D),
                          border_color=RGBColor(0x1A, 0x30, 0x5A))
    add_textbox(slide, pill_x, pill_y + Inches(0.02), pw, pill_h,
                feat, font_size=8, color=ACCENT_LIGHT, bold=False,
                alignment=PP_ALIGN.CENTER)
    pill_x += pw + Inches(0.08)

# ── RIGHT COLUMN: Predictive Heatmap ──
col_right = Inches(6.8)
col_w_r = Inches(5.8)

add_textbox(slide, col_right, col_top, col_w_r, Inches(0.3),
            'AI MODEL 2 \u2014 PREDICTIVE HEATMAP', font_size=9, color=ACCENT, bold=True)

items_right = [
    ('DBSCAN Clustering', 'Groups flagged scans by geographic proximity (~5km radius). Identifies emerging hotspot clusters in real-time.', ACCENT, TEXT_COLOR),
    ('Time Decay', '7-day exponential half-life. A 2-week-old flag counts for 25% intensity. Heatmap always reflects current threat landscape.', PURPLE, TEXT_COLOR),
    ('Trend Prediction', 'Linear regression on weekly cluster snapshots. Extrapolates 2 weeks ahead. Alerts if any cluster is predicted to double in size.', GREEN, TEXT_COLOR),
]

iy = col_top + Inches(0.35)
for title, body, border_col, title_col in items_right:
    add_rec_item(slide, col_right, iy, col_w_r, title, body,
                 border_color=border_col, title_color=title_col)
    iy += Inches(0.92)

# Trust Score label
add_textbox(slide, col_right, iy + Inches(0.05), col_w_r, Inches(0.25),
            'TRUST SCORE \u2014 0 TO 100', font_size=9, color=PURPLE, bold=True)

# Trust score bar
bar_y = iy + Inches(0.35)
bar_h = Inches(0.35)
# Anomaly 45%
w1 = Inches(5.8 * 0.45)
add_rect(slide, col_right, bar_y, w1, bar_h, ACCENT)
add_textbox(slide, col_right, bar_y + Inches(0.04), w1, bar_h,
            'Anomaly 45%', font_size=9, color=PRIMARY, bold=True, alignment=PP_ALIGN.CENTER)
# Heatmap 40%
w2 = Inches(5.8 * 0.40)
add_rect(slide, col_right + w1, bar_y, w2, bar_h, PURPLE)
add_textbox(slide, col_right + w1, bar_y + Inches(0.04), w2, bar_h,
            'Heatmap 40%', font_size=9, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
# Ext 15%
w3 = Inches(5.8 * 0.15)
add_rect(slide, col_right + w1 + w2, bar_y, w3, bar_h, GREEN)
add_textbox(slide, col_right + w1 + w2, bar_y + Inches(0.04), w3, bar_h,
            'Ext 15%', font_size=9, color=PRIMARY, bold=True, alignment=PP_ALIGN.CENTER)

# Threshold boxes
thresholds = [
    ('80\u2013100', 'Safe', GREEN),
    ('50\u201379', 'Warning', YELLOW),
    ('20\u201349', 'Auto-Flag', ORANGE),
    ('0\u201319', 'Critical', RED),
]
th_y = bar_y + Inches(0.5)
th_w = Inches(1.3)
th_h = Inches(0.55)
for i, (val, label, color) in enumerate(thresholds):
    tx = col_right + i * (th_w + Inches(0.13))
    bg_col = RGBColor(
        min(255, color[0] // 6), min(255, color[1] // 6), min(255, color[2] // 6)
    )
    add_rounded_rect(slide, tx, th_y, th_w, th_h,
                     fill_color=bg_col,
                     border_color=RGBColor(
                         min(255, color[0] // 3), min(255, color[1] // 3), min(255, color[2] // 3)
                     ))
    add_textbox(slide, tx, th_y + Inches(0.06), th_w, Inches(0.28),
                val, font_size=14, color=color, alignment=PP_ALIGN.CENTER, font_name='Georgia')
    add_textbox(slide, tx, th_y + Inches(0.32), th_w, Inches(0.2),
                label.upper(), font_size=7, color=color, bold=True, alignment=PP_ALIGN.CENTER)

add_slide_number(slide, 4)


# ────────────────────────── SLIDE 5: INNOVATION + IMPACT ──────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, PRIMARY)
add_section_label(slide, 'What Makes Us Different')
add_section_title(slide, 'Innovation + Impact')
add_body_text(slide, 'Every team here will build QR + blockchain verification. Here\u2019s what we built on top of that.')

# 3 differentiator cards
diffs = [
    ('Semi-Supervised Learning',
     'Our AI trains on normal behavior only. No labeled counterfeit data needed \u2014 which doesn\u2019t exist anyway. Catches anomalies nobody programmed it to look for. Deployable from day one.',
     ACCENT),
    ('Predictive, Not Reactive',
     'Our heatmap doesn\u2019t show where counterfeits were found. It predicts where they\u2019ll appear 2 weeks from now. DBSCAN clustering + linear regression on weekly snapshots = proactive enforcement.',
     PURPLE),
    ('Self-Improving System',
     'Government raids \u2192 confirmed seizure locations \u2192 ground truth training data \u2192 model retrains \u2192 better predictions. Closed feedback loop. Gets smarter with every enforcement action.',
     GREEN),
]
diff_w = Inches(3.6)
diff_h = Inches(1.7)
diff_y = Inches(2.8)
diff_start = Inches(1.2)
diff_gap = Inches(0.4)

for i, (title, body, color) in enumerate(diffs):
    lx = diff_start + i * (diff_w + diff_gap)
    add_card_with_border(slide, lx, diff_y, diff_w, diff_h, title, body,
                         border_side='top', border_color=color)

# 4 impact cards (2x2)
impacts = [
    ('Patient Safety', 'Counterfeit antibiotics drive antimicrobial resistance \u2014 700K+ deaths/year from AMR globally. MediChain gives every patient instant verification before they take a single pill.', RED),
    ('Economic Justice', 'Rural pharmacies bear the blame for counterfeits they unknowingly stock. Real-time verification protects honest businesses and keeps local economies alive.', GREEN),
    ('Public Health Intelligence', 'Every scan adds to India\u2019s first real-time pharmaceutical quality map. Outbreaks caught in days, not months. Government gets data, not quarterly reports.', PURPLE),
    ('Trust in Healthcare', 'When patients doubt their medicine, they stop taking it. Verified drugs restore faith in the healthcare system \u2014 especially in underserved rural communities.', ACCENT),
]
imp_w = Inches(5.5)
imp_h = Inches(1.1)
imp_positions = [
    (Inches(0.8), Inches(4.7)), (Inches(6.8), Inches(4.7)),
    (Inches(0.8), Inches(6.0)), (Inches(6.8), Inches(6.0)),
]
for (title, body, color), (lx, ly) in zip(impacts, imp_positions):
    add_card_with_border(slide, lx, ly, imp_w, imp_h, title, body,
                         border_side='left', border_color=color)

add_slide_number(slide, 5)


# ────────────────────────── SLIDE 6: HOW WE SCALE ──────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, PRIMARY)
add_section_label(slide, 'Feasibility')
add_section_title(slide, 'How We Scale')

# LEFT COLUMN: pricing cards
lx = Inches(0.8)
ly = Inches(1.8)
cw = Inches(5.5)

# Per badge card
ch = Inches(1.2)
bg = add_rounded_rect(slide, lx, ly, cw, ch,
                       fill_color=RGBColor(0x12, 0x1A, 0x2B),
                       border_color=RGBColor(0x1E, 0x29, 0x3B))
add_rect(slide, lx + Emu(Inches(0.05).emu), ly, cw - Emu(Inches(0.1).emu), Pt(3), ACCENT)
add_textbox(slide, lx + Inches(0.3), ly + Inches(0.2), cw - Inches(0.6), Inches(0.3),
            'Per Badge \u2014 \u20B95\u201310', font_size=15, color=TEXT_COLOR, bold=True)
add_textbox(slide, lx + Inches(0.3), ly + Inches(0.55), cw - Inches(0.6), Inches(0.5),
            'Manufacturer pays per medicine unit. Covers QR generation + blockchain registration. One-time cost at production.',
            font_size=10, color=TEXT_MUTED)

ly2 = ly + ch + Inches(0.2)
bg = add_rounded_rect(slide, lx, ly2, cw, ch,
                       fill_color=RGBColor(0x12, 0x1A, 0x2B),
                       border_color=RGBColor(0x1E, 0x29, 0x3B))
add_rect(slide, lx + Emu(Inches(0.05).emu), ly2, cw - Emu(Inches(0.1).emu), Pt(3), PURPLE)
add_textbox(slide, lx + Inches(0.3), ly2 + Inches(0.2), cw - Inches(0.6), Inches(0.3),
            'Per Scan', font_size=15, color=TEXT_COLOR, bold=True)
add_textbox(slide, lx + Inches(0.3), ly2 + Inches(0.55), cw - Inches(0.6), Inches(0.5),
            'Verification fee per scan. India dispenses 30B drug units/year. Even 0.1% adoption = 30M scans/year.',
            font_size=10, color=TEXT_MUTED)

# Info items
ly3 = ly2 + ch + Inches(0.25)
item_h = Inches(0.8)
add_rec_item(slide, lx, ly3, cw,
             'QR codes are already familiar',
             '95% smartphone penetration in urban India. UPI payments made QR scanning second nature.',
             border_color=ACCENT, title_color=TEXT_COLOR)
add_rec_item(slide, lx, ly3 + Inches(0.88), cw,
             'Zero friction for manufacturers',
             'Blister pack printing machines already exist (Domino, Videojet). Same surface already printed with batch/expiry.',
             border_color=GREEN, title_color=TEXT_COLOR)

# RIGHT COLUMN: chart area + callout
rx = Inches(6.8)
rw = Inches(5.8)
chart_h = Inches(2.8)

# Chart placeholder (bar chart representation)
bg = add_rounded_rect(slide, rx, Inches(1.8), rw, chart_h,
                       fill_color=RGBColor(0x12, 0x1A, 0x2B),
                       border_color=RGBColor(0x1E, 0x29, 0x3B))
add_textbox(slide, rx, Inches(1.9), rw, Inches(0.3),
            'Cost Per Unit Comparison (INR)', font_size=11, color=TEXT_COLOR,
            bold=False, alignment=PP_ALIGN.CENTER, font_name='Georgia')

# Bar chart simulation
bar_base_y = Inches(4.2)
bar_max_h = Inches(1.8)
# MediChain bar (short - ~₹7.5)
bar1_h = Inches(0.18)
bar1_w = Inches(1.2)
bar1_x = rx + Inches(1.2)
add_rect(slide, bar1_x, bar_base_y - bar1_h, bar1_w, bar1_h, GREEN)
add_textbox(slide, bar1_x - Inches(0.3), bar_base_y + Inches(0.05), Inches(1.8), Inches(0.25),
            'MediChain', font_size=9, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)
add_textbox(slide, bar1_x + bar1_w + Inches(0.1), bar_base_y - bar1_h, Inches(0.6), Inches(0.2),
            '\u20B97.5', font_size=9, color=GREEN, bold=True)

# Paper bar (tall - ~₹75)
bar2_h = Inches(1.8)
bar2_w = Inches(1.2)
bar2_x = rx + Inches(3.4)
add_rect(slide, bar2_x, bar_base_y - bar2_h, bar2_w, bar2_h, RED)
add_textbox(slide, bar2_x - Inches(0.5), bar_base_y + Inches(0.05), Inches(2.2), Inches(0.25),
            'Paper Traceability', font_size=9, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)
add_textbox(slide, bar2_x + bar2_w + Inches(0.1), bar_base_y - bar2_h, Inches(0.6), Inches(0.2),
            '\u20B975', font_size=9, color=RED, bold=True)

# Bottom callout
callout_y = Inches(4.9)
callout_h = Inches(1.1)
bg = add_rounded_rect(slide, rx, callout_y, rw, callout_h,
                       fill_color=RGBColor(0x10, 0x15, 0x2B),
                       border_color=RGBColor(0x1E, 0x2A, 0x4B))
add_textbox(slide, rx + Inches(0.3), callout_y + Inches(0.15), rw - Inches(0.6), Inches(0.4),
            'We don\u2019t sell QR codes. We sell trust.', font_size=14,
            color=TEXT_COLOR, alignment=PP_ALIGN.CENTER)
add_textbox(slide, rx + Inches(0.3), callout_y + Inches(0.55), rw - Inches(0.6), Inches(0.4),
            'Every scan is proof that the drug is real, living permanently on the blockchain.',
            font_size=10, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

# Footer
add_textbox(slide, Inches(0), Inches(6.8), SLIDE_W, Inches(0.4),
            'Team Code Diggers \u2014 MediChain', font_size=11, color=TEXT_MUTED,
            bold=False, alignment=PP_ALIGN.CENTER)

add_slide_number(slide, 6)


# ────────────────────────── SLIDE 7: TIMELINE ──────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, PRIMARY)
add_section_label(slide, 'Implementation')
add_section_title(slide, '48-Hour Build Timeline')
add_body_text(slide, 'Three-day hackathon sprint from concept to deployed prototype. Feb 13\u201315, 2026.')

# Timeline cards - Row 1
row1 = [
    ('Feb 13 \u2022 AM', 'Scaffolding & Contract', 'Monorepo setup, PharmaChain.sol \u2014 registration, custody, verification, flagging.', 'd1'),
    ('Feb 13 \u2022 Eve', 'Dashboard Frontend', 'Next.js 14 + Tailwind. 6 pages: Overview, Drugs, Heatmap, Alerts, Analytics, Register.', 'd1'),
    ('Feb 14 \u2022 PM', 'Mock Data & Polish', 'Full demo mode, animated SVGs, count-up animations, skeleton loading.', 'd2'),
    ('Feb 15 \u2022 AM', 'Integration & Testing', 'E2E flow testing, contract deploy to Sepolia ETH, demo prep.', 'd3'),
]
row2 = [
    ('Feb 13 \u2022 PM', 'AI & ML Engine', 'FastAPI with Isolation Forest, Autoencoder & DBSCAN heatmap. 60K synthetic records.', 'd1'),
    ('Feb 14 \u2022 AM', 'Flutter Mobile App', 'Consumer & commuter flows, QR scanning, blockchain verify, Riverpod.', 'd2'),
    ('Feb 14 \u2022 Eve', 'Landing Page & Auth', '9 animated SVG components, marketing site, sign-in & sign-up.', 'd2'),
    ('Feb 15 \u2022 PM', 'Final Demo & Submit', 'Live demo: register \u2192 scan \u2192 verify \u2192 flag \u2192 alert. Deck done.', 'd3'),
]

badge_colors = {
    'd1': (RGBColor(0x0E, 0x17, 0x2D), ACCENT_LIGHT),
    'd2': (RGBColor(0x18, 0x10, 0x2D), PURPLE_LIGHT),
    'd3': (RGBColor(0x08, 0x1E, 0x18), GREEN),
}

card_w = Inches(2.8)
card_h = Inches(1.3)
start_x = Inches(0.6)
gap_x = Inches(0.2)

def draw_timeline_row(slide, items, y_pos):
    for i, (badge_text, title, body, day_class) in enumerate(items):
        lx = start_x + i * (card_w + gap_x)
        bg = add_rounded_rect(slide, lx, y_pos, card_w, card_h,
                               fill_color=SURFACE,
                               border_color=RGBColor(0x2A, 0x35, 0x4B))
        # Badge
        bg_col, text_col = badge_colors[day_class]
        badge_w = Inches(1.4)
        badge_h = Inches(0.22)
        badge = add_rounded_rect(slide, lx + Inches(0.15), y_pos + Inches(0.12),
                                 badge_w, badge_h, fill_color=bg_col)
        add_textbox(slide, lx + Inches(0.15), y_pos + Inches(0.1), badge_w, badge_h,
                    badge_text.upper(), font_size=7, color=text_col, bold=True,
                    alignment=PP_ALIGN.CENTER)
        # Title
        add_textbox(slide, lx + Inches(0.15), y_pos + Inches(0.4), card_w - Inches(0.3), Inches(0.25),
                    title, font_size=11, color=TEXT_COLOR, bold=True)
        # Body
        add_textbox(slide, lx + Inches(0.15), y_pos + Inches(0.65), card_w - Inches(0.3), Inches(0.6),
                    body, font_size=8, color=TEXT_MUTED)

draw_timeline_row(slide, row1, Inches(2.7))

# Timeline track line
track_y = Inches(4.15)
line_w = Inches(12)
line_x = Inches(0.6)
# Three-segment colored line
seg_w_emu = int(line_w.emu / 3)
add_rect(slide, line_x, track_y, Emu(seg_w_emu), Pt(3), ACCENT)
add_rect(slide, Emu(line_x.emu + seg_w_emu), track_y, Emu(seg_w_emu), Pt(3), PURPLE)
add_rect(slide, Emu(line_x.emu + seg_w_emu * 2), track_y, Emu(seg_w_emu), Pt(3), GREEN)

# Dots on the line
dot_size = Inches(0.14)
dot_positions = [0, 1, 2, 3, 4, 5, 6, 7]
dot_colors = [ACCENT, ACCENT, ACCENT, PURPLE, PURPLE, PURPLE, GREEN, GREEN]
for i, (pos, dcolor) in enumerate(zip(dot_positions, dot_colors)):
    dx = Emu(line_x.emu + int(line_w.emu * pos / 7) - int(dot_size.emu / 2))
    circle = add_circle(slide, dx, Emu(track_y.emu - int(dot_size.emu / 2) + Pt(1.5).emu), dot_size, dcolor)

draw_timeline_row(slide, row2, Inches(4.5))

# Day labels
day_y = Inches(6.0)
day_labels = [
    ('\u2014\u2014 Day 1 \u2014 Feb 13', ACCENT),
    ('\u2014\u2014 Day 2 \u2014 Feb 14', PURPLE),
    ('\u2014\u2014 Day 3 \u2014 Feb 15', GREEN),
]
for i, (label, color) in enumerate(day_labels):
    dx = Inches(0.8) + i * Inches(4.0)
    add_textbox(slide, dx, day_y, Inches(3), Inches(0.3),
                label, font_size=9, color=color, bold=True)

add_slide_number(slide, 7)


# ══════════════════════════════════════════════════════════════
# Save
# ══════════════════════════════════════════════════════════════
output_path = '/home/sinha/Hackathon/MediChain_Presentation.pptx'
prs.save(output_path)
print(f'Saved to {output_path}')
